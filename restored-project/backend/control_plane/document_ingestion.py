from __future__ import annotations
from pathlib import Path
import csv,json,importlib.util
class DocumentIngestor:
    def __init__(self,knowledge):self.knowledge=knowledge
    def extract(self,path):
        p=Path(path);s=p.suffix.casefold()
        if s in {".txt",".md",".py",".js",".ts",".json",".csv",".xml",".html"}:return p.read_text(encoding="utf-8",errors="replace")
        if s==".pdf" and importlib.util.find_spec("pypdf"):
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
        if s==".docx" and importlib.util.find_spec("docx"):
            from docx import Document
            return "\n".join(x.text for x in Document(str(p)).paragraphs)
        if s==".xlsx" and importlib.util.find_spec("openpyxl"):
            import openpyxl
            wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
            return "\n".join(",".join("" if v is None else str(v) for v in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        if s==".pptx" and importlib.util.find_spec("pptx"):
            from pptx import Presentation
            return "\n".join(shape.text for slide in Presentation(str(p)).slides for shape in slide.shapes if hasattr(shape,"text"))
        raise RuntimeError(f"unsupported document type or missing parser: {p.suffix}")
    def ingest(self,path):return self.knowledge.ingest_text(str(Path(path)),self.extract(path))
