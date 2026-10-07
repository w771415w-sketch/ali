# -*- coding: utf-8 -*-
"""Real parsers for documents and archives. Optional heavy readers degrade gracefully."""
from __future__ import annotations
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional
import csv, io, json, tarfile, zipfile, gzip, bz2, lzma, xml.etree.ElementTree as ET
import re

@dataclass
class ParsedDocument:
    path: str
    kind: str
    text: str
    metadata: Dict[str, Any]
    tables: List[List[List[str]]]
    extracted_files: List[str]
    warning: Optional[str] = None
    def to_dict(self): return asdict(self)

def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ('utf-8','utf-8-sig','utf-16','cp1256','cp1252','latin-1'):
        try: return data.decode(enc)
        except Exception: pass
    return data.decode('utf-8','replace')

def _parse_json(path: Path) -> str:
    obj = json.loads(_read_text(path))
    return json.dumps(obj, ensure_ascii=False, indent=2)

def _parse_csv(path: Path) -> tuple[str, List[List[List[str]]]]:
    text = _read_text(path)
    rows = list(csv.reader(io.StringIO(text)))
    return '\n'.join(' | '.join(r) for r in rows), [rows]

def _parse_pdf(path: Path) -> tuple[str, List[List[List[str]]], Dict[str,Any], Optional[str]]:
    tables = []; meta={}; warnings=[]
    try:
        import fitz
        doc = fitz.open(str(path))
        parts=[]
        meta.update({"page_count": doc.page_count, "metadata": doc.metadata})
        for page in doc:
            parts.append(page.get_text("text"))
            try:
                tf = page.find_tables()
                for table in tf.tables:
                    tables.append(table.extract())
            except Exception:
                pass
        text='\n\n'.join(parts)
        # OCR only when the PDF appears scanned and a real Tesseract stack is available.
        if len(text.strip()) < max(40, doc.page_count * 20):
            try:
                import pytesseract
                from PIL import Image
                ocr_parts=[]
                for page in doc:
                    pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5), alpha=False)
                    img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
                    ocr_parts.append(pytesseract.image_to_string(img,lang='ara+eng'))
                ocr_text='\n\n'.join(ocr_parts).strip()
                if ocr_text:
                    text=ocr_text; meta['ocr']=True
                else: warnings.append('PDF scan detected but OCR returned no text')
            except Exception:
                warnings.append('PDF has little/no selectable text; install Tesseract + pytesseract for OCR')
        return text,tables,meta,'; '.join(warnings) or None
    except Exception:
        try:
            from pypdf import PdfReader
            reader=PdfReader(str(path))
            parts=[p.extract_text() or '' for p in reader.pages]
            return '\n\n'.join(parts),tables,{"page_count":len(reader.pages)},None
        except Exception as e:
            return '',[],{},f'PDF parser unavailable: {e}'

def _parse_docx(path: Path) -> tuple[str,List[List[List[str]]],Dict[str,Any]]:
    from docx import Document
    doc=Document(str(path)); parts=[p.text for p in doc.paragraphs if p.text.strip()]; tables=[]
    for t in doc.tables:
        rows=[]
        for r in t.rows: rows.append([c.text for c in r.cells])
        tables.append(rows)
    return '\n'.join(parts), tables, {"paragraphs":len(doc.paragraphs),"tables":len(doc.tables)}

def _parse_html(path: Path) -> str:
    from bs4 import BeautifulSoup
    soup=BeautifulSoup(_read_text(path),'html.parser')
    for tag in soup(['script','style','noscript']): tag.decompose()
    return soup.get_text('\n')

def _xml_text(path: Path) -> str:
    root=ET.fromstring(_read_text(path)); return '\n'.join(t.strip() for t in root.itertext() if t.strip())

def safe_extract_zip(path: Path, dest: Path) -> List[str]:
    out=[]; dest=dest.resolve(); dest.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            target=(dest/info.filename).resolve()
            if dest not in target.parents and target != dest: raise ValueError(f'archive path traversal blocked: {info.filename}')
        for info in z.infolist():
            if info.is_dir(): continue
            target=(dest/info.filename).resolve(); target.parent.mkdir(parents=True,exist_ok=True)
            with z.open(info) as src, open(target,'wb') as dst: dst.write(src.read())
            out.append(str(target))
    return out

def safe_extract_tar(path: Path, dest: Path) -> List[str]:
    out=[]; dest=dest.resolve(); dest.mkdir(parents=True,exist_ok=True)
    with tarfile.open(path) as t:
        members=[]
        for m in t.getmembers():
            target=(dest/m.name).resolve()
            if dest not in target.parents and target != dest: raise ValueError(f'archive path traversal blocked: {m.name}')
            if m.isdir() or m.isfile(): members.append(m)
        t.extractall(dest, members=members)
        out=[str((dest/m.name).resolve()) for m in members if m.isfile()]
    return out


def safe_extract_optional(path: Path, dest: Path) -> List[str]:
    ext=path.suffix.lower(); dest=dest.resolve(); dest.mkdir(parents=True,exist_ok=True)
    if ext=='.rar':
        try:
            import rarfile
            with rarfile.RarFile(path) as rf:
                for info in rf.infolist():
                    target=(dest/info.filename).resolve()
                    if dest not in target.parents and target!=dest: raise ValueError(f'archive path traversal blocked: {info.filename}')
                rf.extractall(dest)
            return [str(p) for p in dest.rglob('*') if p.is_file()]
        except ImportError:
            import shutil, subprocess
            exe=shutil.which('7z') or shutil.which('7zz') or shutil.which('7za') or shutil.which('unrar')
            if not exe: raise RuntimeError('RAR support needs rarfile or a 7-Zip/unrar executable on PATH')
            subprocess.run([exe,'x','-y',str(path),f'-o{dest}'],check=True,capture_output=True,text=True,timeout=600)
            return [str(p) for p in dest.rglob('*') if p.is_file()]
    if ext=='.7z':
        try:
            import py7zr
            with py7zr.SevenZipFile(path,'r') as z: z.extractall(dest)
            return [str(p) for p in dest.rglob('*') if p.is_file()]
        except ImportError:
            import shutil, subprocess
            exe=shutil.which('7z') or shutil.which('7zz') or shutil.which('7za')
            if not exe: raise RuntimeError('7z support needs py7zr or a 7-Zip executable on PATH')
            subprocess.run([exe,'x','-y',str(path),f'-o{dest}'],check=True,capture_output=True,text=True,timeout=600)
            return [str(p) for p in dest.rglob('*') if p.is_file()]
    raise RuntimeError(f'unsupported archive: {path.suffix}')

def parse_file(path: str|Path, extract_root: str|Path|None=None) -> ParsedDocument:
    p=Path(path); ext=p.suffix.lower(); kind=ext.lstrip('.') or 'unknown'
    metadata={"name":p.name,"size":p.stat().st_size}
    text=''; tables=[]; extracted=[]; warning=None
    try:
        if ext in {'.txt','.md','.py','.js','.ts','.tsx','.jsx','.css','.jsonl','.log','.ini','.yaml','.yml','.toml'}:
            text=_read_text(p)
        elif ext=='.json': text=_parse_json(p); kind='json'
        elif ext=='.csv': text,tables=_parse_csv(p); kind='csv'
        elif ext=='.xml': text=_xml_text(p); kind='xml'
        elif ext in {'.html','.htm'}: text=_parse_html(p); kind='html'
        elif ext=='.pdf': text,tables,meta2,warning=_parse_pdf(p); metadata.update(meta2); kind='pdf'
        elif ext=='.docx': text,tables,meta2=_parse_docx(p); metadata.update(meta2); kind='docx'
        elif ext=='.zip' and extract_root:
            extracted=safe_extract_zip(p,Path(extract_root)/p.stem); text='\n'.join(extracted); kind='archive'
        elif ext in {'.rar','.7z'} and extract_root:
            extracted=safe_extract_optional(p,Path(extract_root)/p.stem); text='\n'.join(extracted); kind='archive'
        elif ext in {'.tar','.gz','.bz2','.xz','.tgz','.tbz2'} and extract_root:
            if tarfile.is_tarfile(p): extracted=safe_extract_tar(p,Path(extract_root)/p.stem); kind='archive'; text='\n'.join(extracted)
            elif ext=='.gz': text=gzip.decompress(p.read_bytes()).decode('utf-8','replace')
            elif ext=='.bz2': text=bz2.decompress(p.read_bytes()).decode('utf-8','replace')
            elif ext=='.xz': text=lzma.decompress(p.read_bytes()).decode('utf-8','replace')
        else:
            warning='Unsupported or binary format; file indexed by metadata only'
    except Exception as e:
        warning=str(e)
    return ParsedDocument(str(p),kind,text,metadata,tables,extracted,warning)
